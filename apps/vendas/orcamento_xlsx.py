"""
Planilha Excel INTERNA do orçamento (para discussão com a gerência).

Abas: Resumo | MOB | Custos Diretos | Impostos e Markup | Resultado Detalhado
As abas de entrada têm fórmulas vivas: alterar quantidade, prazo, valor,
percentuais ou markup recalcula tudo, inclusive o Valor de Venda.
Mostra margem e lucro — NÃO enviar ao cliente.
"""
import io

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

from .orcamento_dados import montar_dados_orcamento

AZUL = '1E40AF'
AZUL_CLARO = 'DBEAFE'
CINZA = 'F1F5F9'
VERDE = '16A34A'
LARANJA = 'D97706'
BRL = '"R$" #,##0.00'
PCT = '0.00"%"'

_borda = Side(style='thin', color='D1D5DB')
BORDA = Border(left=_borda, right=_borda, top=_borda, bottom=_borda)


def _titulo(ws, texto, sub, largura):
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=largura)
    c = ws.cell(row=1, column=1, value=texto)
    c.font = Font(name='Calibri', size=16, bold=True, color='FFFFFF')
    c.fill = PatternFill('solid', fgColor=AZUL)
    c.alignment = Alignment(vertical='center', indent=1)
    ws.row_dimensions[1].height = 30
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=largura)
    s = ws.cell(row=2, column=1, value=sub)
    s.font = Font(italic=True, size=10, color='6B7280')
    s.alignment = Alignment(indent=1)


def _cabecalho(ws, linha, titulos):
    for i, t in enumerate(titulos, start=1):
        c = ws.cell(row=linha, column=i, value=t)
        c.font = Font(bold=True, color='FFFFFF', size=11)
        c.fill = PatternFill('solid', fgColor=AZUL)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = BORDA
    ws.row_dimensions[linha].height = 24


def _larguras(ws, larguras):
    for i, w in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _celula(ws, ref, valor, fmt=None, negrito=False, cor=None, fundo=None, alinh=None):
    c = ws[ref]
    c.value = valor
    c.border = BORDA
    c.font = Font(bold=negrito, color=cor or '111827')
    if fmt:
        c.number_format = fmt
    if fundo:
        c.fill = PatternFill('solid', fgColor=fundo)
    if alinh:
        c.alignment = Alignment(horizontal=alinh)
    return c


def gerar_planilha_orcamento(proposta):
    d = montar_dados_orcamento(proposta)
    sub = f'{proposta.codigo} — {proposta.titulo}'.strip(' —')
    cliente = proposta.get_cliente_nome()
    if cliente:
        sub += f'  |  Cliente: {cliente}'
    sub += '  |  USO INTERNO — contém margem e lucro'

    wb = Workbook()
    ws_res = wb.active
    ws_res.title = 'Resumo'
    ws_mob = wb.create_sheet('MOB')
    ws_cd = wb.create_sheet('Custos Diretos')
    ws_pc = wb.create_sheet('Impostos e Markup')
    ws_det = wb.create_sheet('Resultado Detalhado')

    # ── MOB ─────────────────────────────────────────────────────────
    _titulo(ws_mob, 'Mão de Obra (MOB)', sub, 5)
    _cabecalho(ws_mob, 4, ['Descrição', 'Qtd', 'Prazo (meses)', 'Valor Unit. (R$)', 'Valor Total (R$)'])
    _larguras(ws_mob, [52, 12, 16, 20, 20])
    r = 5
    for it in d['mob']:
        _celula(ws_mob, f'A{r}', it['descricao'])
        _celula(ws_mob, f'B{r}', it['quantidade'], '#,##0.##', cor='1D4ED8')
        _celula(ws_mob, f'C{r}', it['prazo_meses'], '#,##0.##', cor='1D4ED8')
        _celula(ws_mob, f'D{r}', it['valor_unitario'], BRL, cor='1D4ED8')
        _celula(ws_mob, f'E{r}', f'=B{r}*C{r}*D{r}', BRL)
        r += 1
    mob_ini, mob_fim = 5, r - 1
    mob_tot = r
    _celula(ws_mob, f'A{mob_tot}', 'TOTAL MOB', negrito=True, fundo=AZUL_CLARO)
    for col in 'BCD':
        _celula(ws_mob, f'{col}{mob_tot}', None, fundo=AZUL_CLARO)
    _celula(ws_mob, f'E{mob_tot}',
            f'=SUM(E{mob_ini}:E{mob_fim})' if mob_fim >= mob_ini else 0,
            BRL, negrito=True, fundo=AZUL_CLARO)
    ws_mob.freeze_panes = 'A5'

    # ── Custos diretos ──────────────────────────────────────────────
    _titulo(ws_cd, 'Custos Diretos', sub, 5)
    _cabecalho(ws_cd, 4, ['Descrição', 'Qtd', 'Repetição', 'Valor Unit. (R$)', 'Valor Total (R$)'])
    _larguras(ws_cd, [52, 12, 16, 20, 20])
    r = 5
    for it in d['custos_diretos']:
        _celula(ws_cd, f'A{r}', it['descricao'])
        _celula(ws_cd, f'B{r}', it['quantidade'], '#,##0.##', cor='1D4ED8')
        _celula(ws_cd, f'C{r}', it['repeticao'], '#,##0.##', cor='1D4ED8')
        _celula(ws_cd, f'D{r}', it['valor_unitario'], BRL, cor='1D4ED8')
        _celula(ws_cd, f'E{r}', f'=B{r}*C{r}*D{r}', BRL)
        r += 1
    cd_ini, cd_fim = 5, r - 1
    cd_tot = r
    _celula(ws_cd, f'A{cd_tot}', 'TOTAL CUSTOS DIRETOS', negrito=True, fundo=AZUL_CLARO)
    for col in 'BCD':
        _celula(ws_cd, f'{col}{cd_tot}', None, fundo=AZUL_CLARO)
    _celula(ws_cd, f'E{cd_tot}',
            f'=SUM(E{cd_ini}:E{cd_fim})' if cd_fim >= cd_ini else 0,
            BRL, negrito=True, fundo=AZUL_CLARO)
    ws_cd.freeze_panes = 'A5'

    # ── Impostos e markup ───────────────────────────────────────────
    _titulo(ws_pc, 'Impostos, Taxas e Markup (sobre o valor de venda)', sub, 3)
    _cabecalho(ws_pc, 4, ['Item', '% sobre o PV', 'Valor (R$)'])
    _larguras(ws_pc, [44, 18, 22])
    r = 5
    for p in d['percentuais']:
        _celula(ws_pc, f'A{r}', p['nome'])
        _celula(ws_pc, f'B{r}', p['percentual'], PCT, cor='1D4ED8')
        r += 1
    pc_ini, pc_fim = 5, r - 1
    pc_tot = r
    mk_row = r + 1
    # Valores em R$ dependem do PV (célula do Resumo) — preenchidos abaixo.
    _celula(ws_pc, f'A{pc_tot}', 'TOTAL IMPOSTOS E TAXAS', negrito=True, fundo=AZUL_CLARO)
    _celula(ws_pc, f'B{pc_tot}',
            f'=SUM(B{pc_ini}:B{pc_fim})' if pc_fim >= pc_ini else 0,
            PCT, negrito=True, fundo=AZUL_CLARO)
    _celula(ws_pc, f'A{mk_row}', 'MARKUP / LUCRO DESEJADO', negrito=True, fundo='EDE9FE')
    _celula(ws_pc, f'B{mk_row}', d['markup'], PCT, negrito=True, cor='1D4ED8', fundo='EDE9FE')

    # ── Resumo ──────────────────────────────────────────────────────
    _titulo(ws_res, 'Orçamento — Resumo Executivo', sub, 6)
    _larguras(ws_res, [36, 22, 16, 4, 30, 18])
    ws_res['A3'].value = 'Células em azul nas outras abas são editáveis; tudo recalcula sozinho.'
    ws_res['A3'].font = Font(italic=True, size=9, color='6B7280')

    _cabecalho(ws_res, 5, ['Indicador', 'Valor', '% do PV'])
    linhas = [
        (6, 'Custo MOB', f"='MOB'!E{mob_tot}", BRL),
        (7, 'Custos Diretos', f"='Custos Diretos'!E{cd_tot}", BRL),
        (8, 'Custo Base (MOB + Diretos)', '=B6+B7', BRL),
        (9, 'Impostos e taxas (%)', f"='Impostos e Markup'!B{pc_tot}", PCT),
        (10, 'Markup / Lucro (%)', f"='Impostos e Markup'!B{mk_row}", PCT),
        (11, 'VALOR DE VENDA (PV)', '=IF(AND(B8>0,B9+B10<100),B8/(1-(B9+B10)/100),0)', BRL),
        (12, 'Impostos e taxas (R$)', '=B11*B9/100', BRL),
        (13, 'Receita líquida (PV − impostos)', '=B11-B12', BRL),
        (14, 'LUCRO (R$)', '=B11*B10/100', BRL),
        (15, 'Lucro sobre o custo base (%)', '=IF(B8>0,B14/B8*100,0)', PCT),
    ]
    destaque = {11: AZUL_CLARO, 14: 'DCFCE7', 8: CINZA}
    for ln, nome, formula, fmt in linhas:
        neg = ln in destaque
        _celula(ws_res, f'A{ln}', nome, negrito=neg, fundo=destaque.get(ln))
        _celula(ws_res, f'B{ln}', formula, fmt, negrito=neg, fundo=destaque.get(ln),
                cor=VERDE if ln == 14 else None)
        if ln in (6, 7, 8, 12, 13, 14):
            _celula(ws_res, f'C{ln}', f'=IF($B$11>0,B{ln}/$B$11*100,0)', PCT, fundo=destaque.get(ln))
        else:
            _celula(ws_res, f'C{ln}', None, fundo=destaque.get(ln))
    ws_res['A17'].value = '=IF(B9+B10>=100,"ATENÇÃO: impostos + markup somam 100% ou mais — PV inválido.","")'
    ws_res['A17'].font = Font(bold=True, color='DC2626')

    # Tabela de composição do PV (alimenta os gráficos)
    for i, t in enumerate(['Composição do PV', 'Valor (R$)'], start=5):
        c = ws_res.cell(row=5, column=i, value=t)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=AZUL)
        c.alignment = Alignment(horizontal='center')
        c.border = BORDA
    comp = [
        (6, 'Mão de Obra (MOB)', '=B6'),
        (7, 'Custos Diretos', '=B7'),
        (8, 'Impostos e Taxas', '=B12'),
        (9, 'Lucro', '=B14'),
    ]
    for ln, nome, f in comp:
        _celula(ws_res, f'E{ln}', nome)
        _celula(ws_res, f'F{ln}', f, BRL)

    rosca = DoughnutChart()
    rosca.title = 'Composição do valor de venda'
    rosca.holeSize = 55
    rosca.add_data(Reference(ws_res, min_col=6, min_row=5, max_row=9), titles_from_data=True)
    rosca.set_categories(Reference(ws_res, min_col=5, min_row=6, max_row=9))
    rosca.dataLabels = DataLabelList()
    rosca.dataLabels.showPercent = True
    rosca.height, rosca.width = 8.5, 12
    ws_res.add_chart(rosca, 'A19')

    barras = BarChart()
    barras.type = 'col'
    barras.title = 'Custo × Impostos × Lucro × Venda'
    for i, (nome, ref) in enumerate([('Custo Base', 'B8'), ('Impostos', 'B12'), ('Lucro', 'B14'), ('Venda', 'B11')]):
        ws_res.cell(row=30 + i, column=5, value=nome)
        ws_res.cell(row=30 + i, column=6, value=f'={ref}').number_format = BRL
    barras.add_data(Reference(ws_res, min_col=6, min_row=30, max_row=33), titles_from_data=False)
    barras.set_categories(Reference(ws_res, min_col=5, min_row=30, max_row=33))
    barras.legend = None
    barras.height, barras.width = 8.5, 12
    ws_res.add_chart(barras, 'E19')

    # Valores em R$ na aba de impostos (ligados ao PV do Resumo)
    for r in range(pc_ini, pc_fim + 1):
        _celula(ws_pc, f'C{r}', f'=Resumo!$B$11*B{r}/100', BRL)
    _celula(ws_pc, f'C{pc_tot}', f'=SUM(C{pc_ini}:C{pc_fim})' if pc_fim >= pc_ini else 0,
            BRL, negrito=True, fundo=AZUL_CLARO)
    _celula(ws_pc, f'C{mk_row}', f'=Resumo!$B$11*B{mk_row}/100', BRL, negrito=True, fundo='EDE9FE')

    # ── Resultado detalhado (demonstrativo) ─────────────────────────
    _titulo(ws_det, 'Resultado Detalhado do Orçamento', sub, 3)
    _cabecalho(ws_det, 4, ['Demonstrativo', 'Valor (R$)', '% do PV'])
    _larguras(ws_det, [48, 22, 14])
    demo = [
        ('Valor de Venda (PV)', '=Resumo!B11', True, AZUL_CLARO),
    ]
    r = 5
    _celula(ws_det, f'A{r}', demo[0][0], negrito=True, fundo=AZUL_CLARO)
    _celula(ws_det, f'B{r}', demo[0][1], BRL, negrito=True, fundo=AZUL_CLARO)
    _celula(ws_det, f'C{r}', f'=IF($B$5>0,B{r}/$B$5*100,0)', PCT, fundo=AZUL_CLARO)
    r += 1
    for i, p in enumerate(d['percentuais']):
        _celula(ws_det, f'A{r}', f"(−) {p['nome']}")
        _celula(ws_det, f'B{r}', f"=-'Impostos e Markup'!C{pc_ini + i}", BRL)
        _celula(ws_det, f'C{r}', f'=IF($B$5>0,B{r}/$B$5*100,0)', PCT)
        r += 1
    _celula(ws_det, f'A{r}', '(=) Receita líquida', negrito=True, fundo=CINZA)
    _celula(ws_det, f'B{r}', f'=SUM(B5:B{r-1})', BRL, negrito=True, fundo=CINZA)
    _celula(ws_det, f'C{r}', f'=IF($B$5>0,B{r}/$B$5*100,0)', PCT, fundo=CINZA)
    rl = r
    r += 1
    _celula(ws_det, f'A{r}', '(−) Mão de Obra (MOB)')
    _celula(ws_det, f'B{r}', '=-Resumo!B6', BRL)
    _celula(ws_det, f'C{r}', f'=IF($B$5>0,B{r}/$B$5*100,0)', PCT)
    r += 1
    _celula(ws_det, f'A{r}', '(−) Custos Diretos')
    _celula(ws_det, f'B{r}', '=-Resumo!B7', BRL)
    _celula(ws_det, f'C{r}', f'=IF($B$5>0,B{r}/$B$5*100,0)', PCT)
    r += 1
    _celula(ws_det, f'A{r}', '(=) LUCRO', negrito=True, fundo='DCFCE7')
    _celula(ws_det, f'B{r}', f'=B{rl}+B{r-2}+B{r-1}', BRL, negrito=True, cor=VERDE, fundo='DCFCE7')
    _celula(ws_det, f'C{r}', f'=IF($B$5>0,B{r}/$B$5*100,0)', PCT, negrito=True, fundo='DCFCE7')
    r += 2
    ws_det[f'A{r}'].value = 'Conferência: lucro do demonstrativo − lucro do Resumo (deve ser 0,00)'
    ws_det[f'A{r}'].font = Font(italic=True, size=9, color='6B7280')
    _celula(ws_det, f'B{r}', f'=ROUND(B{r-2}-Resumo!B14,2)', BRL)

    # Impressão: A4 paisagem, ajustado à largura.
    for ws in wb.worksheets:
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
